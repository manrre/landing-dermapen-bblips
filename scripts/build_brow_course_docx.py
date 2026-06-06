from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageEnhance, ImageFilter


WORKSPACE = Path(r"C:\Users\Usuario\Documents\New project")
EXTRACT_DIR = WORKSPACE / "tmp_pdf_extract"
OUTPUT_DIR = WORKSPACE / "artifacts"
OUTPUT_DIR.mkdir(exist_ok=True)
ENHANCED_DIR = OUTPUT_DIR / "enhanced_images"
ENHANCED_DIR.mkdir(exist_ok=True)

RAW_DOCX = OUTPUT_DIR / "diseno_de_cejas_reestructurado_raw.docx"
WATERMARKED_DOCX = OUTPUT_DIR / "diseno_de_cejas_reestructurado_editable.docx"
FINAL_RAW_DOCX = OUTPUT_DIR / "diseno_de_cejas_manual_final_raw.docx"
FINAL_DOCX = OUTPUT_DIR / "diseno_de_cejas_manual_final.docx"
PROMPTS_MD = OUTPUT_DIR / "image_regeneration_prompts.md"
SKIP_IMAGES = {
    "page_01_img_01.png",
    "page_01_img_02.png",
}

VIDEO_LINKS = {
    13: [
        ("Video: diseno de cejas", "https://youtu.be/9fFBVRmSw88?si=aTozANbaZ5eQizYP"),
    ],
    17: [
        ("Video: depilacion de cejas", "https://youtu.be/r7B6u5o76SE?si=rtpwMn_t0g27RbG1"),
    ],
    18: [
        ("Video: visagismo de cejas", "https://youtu.be/_tqkl2SQGi4?si=cKsUjX_oZnZ3OXVr"),
    ],
    22: [
        ("Video: procedimiento de cejas con henna", "https://youtu.be/tTU4nYJwWbY?si=4gtSVPysqNI4HLbs"),
    ],
    25: [
        ("Video: laminado de cejas", "https://youtu.be/xIAzqFuKokQ?si=RuUI_ZuNOmKqUAMc"),
    ],
    41: [
        ("Video: profundidad en micropigmentacion", "https://youtu.be/pjeJPidc4Gg?si=pIMDdtFBBRKErzBm"),
    ],
    42: [
        ("Video: cejas pelo a pelo", "https://youtu.be/MgEcSr5g6Uc?si=0T-Qff4sc_BwSSUe"),
    ],
    44: [
        ("Video: cejas mixtas", "https://youtu.be/502110zannM?si=SA4jjNxlhjK_vIAt"),
    ],
}

IMAGE_PROMPTS = {
    "page_01_img_03.jpg": "Close-up frontal beauty portrait of a woman's face with professional eyebrow mapping lines and yellow measurement points, clean studio lighting, symmetrical composition, educational brow design guide, realistic skin texture, no watermark.",
    "page_08_img_01.jpg": "Close-up of a cosmetologist gently cleaning an eyebrow area with a cotton pad, soft clinical beauty lighting, clean skin-care training manual style, high-resolution photo, no text.",
    "page_10_img_01.png": "Educational eyebrow measurement practice sheet showing a clean eye drawing, brow mapping angles, labeled points A B C and proportional measurement notes, crisp printable training diagram, white background.",
    "page_15_img_01.jpg": "Clean infographic comparing eyebrow hair removal methods: tweezers, thread, wax, blade and electric trimmer, minimal line icons, Spanish labels, salon training manual style.",
    "page_24_img_01.png": "Brow lamination educational illustration with labeled eyebrow hair direction, before and after grooming effect, pastel but professional beauty academy style, crisp text-free layout.",
    "page_35_img_01.jpg": "Scientific educational diagram of skin layers for micropigmentation training: epidermis, dermis, hypodermis, simplified cross-section, clean labels in Spanish, medical beauty manual style.",
    "page_41_img_01.png": "Macro close-up of a professionally microbladed eyebrow showing fine hair-like strokes, natural brown pigment, high detail, neutral background, no watermark.",
    "page_48_img_01.png": "Aftercare icons for eyebrow micropigmentation: avoid water, avoid sweat, avoid sun, avoid makeup, do not scratch, clean black line icons on white background, professional training sheet.",
    "page_51_img_01.jpg": "Eyebrow healing timeline collage showing day 1 to day 6 changes after micropigmentation, clean clinical layout, consistent lighting, high-resolution educational reference.",
    "page_56_img_02.jpg": "Detailed product-style image of a 7F microblading needle cartridge, isolated on white background, crisp technical beauty training asset, no logo.",
    "page_57_img_01.jpg": "Set of microblading blade tip shapes arranged in a row with small technical labels, clean white background, high-resolution educational tool diagram.",
}

LICENSE_RE = re.compile(
    r"Licensed to Beauty Trends\s*-\s*beautytrendscol@gmail\.com\s*-\s*HP3930942771",
    re.IGNORECASE,
)

TEXT_FIXES = {
    "turostro": "tu rostro",
    "dela": "de la",
    "delas": "de las",
    "delos": "de los",
    "decejas": "de cejas",
    "rostrómás": "rostro más",
    "rostromás": "rostro más",
    "parezcamás": "parezca más",
    "puedenhacer": "pueden hacer",
    "hayalgunas": "hay algunas",
    "naturaldel": "natural del",
    "rostrocuadrado": "rostro cuadrado",
    "formacuadrada": "forma cuadrada",
    "áreasespecíficas": "áreas específicas",
    "unasemana": "una semana",
    "primerosdías": "primeros días",
    "aplicaciónde": "aplicación de",
    "cuidadoposterior": "cuidado posterior",
    "cicatrizaciónadecuada": "cicatrización adecuada",
    "preparacion": "preparación",
    "limpiado": "limpieza",
    "depilaciónpuede": "depilación puede",
    "técnicascomo": "técnicas como",
    "técnicadependerá": "técnica dependerá",
    "preferenciasindividuales": "preferencias individuales",
    "omicroblading": "o microblading",
    "micropigmentacion": "micropigmentación",
    "microbranding": "microblading",
    "laparte": "la parte",
    "lashoras": "las horas",
    "lospdf": "los PDF",
    "prácticausando": "práctica usando",
    "practicas másalcanzas": "practicas, más alcanzas",
    "clientesresultados": "clientes resultados",
    "bueno.suerte": "bueno. Suerte",
    "compartirtus": "compartir tus",
    "grupode": "grupo de",
    ".bejos": ". Besos",
}

SECTIONS = [
    {
        "title": "Portada y Mapa Del Curso",
        "summary": "Apertura general y listado de modulos originales.",
        "pages": [],
    },
    {
        "title": "Fundamentos Del Diseno De Cejas",
        "summary": "Forma del rostro, preparacion, limpieza y medicion inicial.",
        "pages": [4, 5, 6, 7, 8, 9, 10, 11, 12, 13],
    },
    {
        "title": "Tipos De Cejas y Depilacion",
        "summary": "Clasificacion de cejas, metodos de depilacion y visagismo.",
        "pages": [14, 15, 16, 17, 18, 19, 20],
    },
    {
        "title": "Procedimientos Con Henna y Laminado",
        "summary": "Aplicacion, materiales, pasos y duracion de resultados.",
        "pages": [21, 22, 23, 24, 25, 26],
    },
    {
        "title": "Micropigmentacion: Bases y Seguridad",
        "summary": "Introduccion, bioseguridad, anestesia, pigmentos y criterios previos.",
        "pages": [27, 28, 29, 30, 31, 32, 33, 34],
    },
    {
        "title": "Profundidad De La Piel y Tecnica",
        "summary": "Capas de la piel, profundidad del microblading y tecnica pelo a pelo.",
        "pages": [35, 36, 37, 38, 39, 40, 41, 42],
    },
    {
        "title": "Cejas Mixtas y Sombreadas",
        "summary": "Tecnicas hibridas, sombreado y herramientas relacionadas.",
        "pages": [43, 44, 45, 46, 47],
    },
    {
        "title": "Cuidados, Cicatrizacion y Agujas",
        "summary": "Postprocedimiento, evolucion de la cicatrizacion y tipologia de agujas.",
        "pages": [48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59],
    },
    {
        "title": "Cierre y Recomendaciones Finales",
        "summary": "Aplicacion practica y recomendacion final del curso original.",
        "pages": [60],
    },
]


def clean_text(text: str) -> str:
    text = LICENSE_RE.sub("", text)
    text = text.replace("\uf0b7", "-")
    text = text.replace("\xa0", " ")
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    for wrong, right in TEXT_FIXES.items():
        cleaned = cleaned.replace(wrong, right)
    cleaned = re.sub(r"(?<=[a-záéíóúñ])(?=[A-ZÁÉÍÓÚÑ])", " ", cleaned)
    return cleaned.strip()


def paragraph_chunks(text: str) -> list[str]:
    parts = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    if len(parts) > 1:
        return parts
    text = text.replace(" . ", ". ")
    text = re.sub(r"(?<=[\.:;])(?=[A-ZÁÉÍÓÚÑ¿])", "\n", text)
    text = re.sub(r"(?<=\d\.) (?=[A-ZÁÉÍÓÚÑ])", "\n", text)
    return [p.strip() for p in text.split("\n") if p.strip()]


def page_heading(page_num: int, text: str) -> str:
    first_line = text.splitlines()[0].strip() if text.splitlines() else ""
    first_line = re.sub(r"\s+", " ", first_line)
    if first_line and len(first_line) <= 90:
        return first_line
    if ":" in first_line:
        title = first_line.split(":", 1)[0].strip()
        if title and len(title) <= 70:
            return title + ":"
    if first_line:
        return first_line[:70].rstrip(" ,.;") + "..."
    return f"Continuacion del tema {page_num}"


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5E8C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    for style_name, size, color, bold in [
        ("Title", 24, RGBColor(41, 67, 94), True),
        ("Heading 1", 18, RGBColor(85, 55, 31), True),
        ("Heading 2", 13, RGBColor(41, 67, 94), True),
        ("Heading 3", 11, RGBColor(96, 96, 96), True),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = bold


def setup_section(section) -> None:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    header = section.header
    if header.paragraphs:
        p = header.paragraphs[0]
    else:
        p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("@formacion.integral.artisca")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 120, 120)

    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = ""


def upscale_for_docx(image_path: Path) -> BytesIO:
    source = enhanced_image_path(image_path)
    with Image.open(source) as img:
        img = img.convert("RGB")
        bio = BytesIO()
        img.save(bio, format="PNG")
        bio.seek(0)
        return bio


def enhanced_image_path(image_path: Path) -> Path:
    out = ENHANCED_DIR / f"{image_path.stem}_enhanced.png"
    if out.exists():
        return out

    with Image.open(image_path) as img:
        img = img.convert("RGB")
        target_width = 1400 if img.width < 700 else img.width
        if img.width < target_width:
            ratio = target_width / max(img.width, 1)
            img = img.resize((target_width, int(img.height * ratio)), Image.Resampling.LANCZOS)
        elif img.width > 1800:
            ratio = 1800 / img.width
            img = img.resize((1800, int(img.height * ratio)), Image.Resampling.LANCZOS)

        img = ImageEnhance.Contrast(img).enhance(1.06)
        img = ImageEnhance.Sharpness(img).enhance(1.35)
        img = img.filter(ImageFilter.UnsharpMask(radius=1.2, percent=115, threshold=3))
        img.save(out, "PNG", optimize=True)
    return out


def add_image_block(doc: Document, image_path: Path, caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(upscale_for_docx(image_path), width=Inches(5.45))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = RGBColor(110, 110, 110)


def iter_page_images(page_num: int) -> list[Path]:
    paths = []
    for path in sorted(EXTRACT_DIR.glob(f"page_{page_num:02d}_img_*.*")):
        if path.name in SKIP_IMAGES:
            continue
        paths.append(path)
    return paths


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Diseno Profesional De Cejas")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(84, 54, 35)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Guia practica de diseno, visagismo, henna, laminado y micropigmentacion")
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(70, 70, 70)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("Manual de formacion para consulta, estudio y practica profesional.")
    nr.font.size = Pt(10.5)

    for img_path in iter_page_images(1):
        add_image_block(doc, img_path)

    doc.add_page_break()


def add_module_map(doc: Document) -> None:
    doc.add_heading("Mapa Del Documento", level=1)
    intro = doc.add_paragraph(
        "La siguiente organizacion presenta el recorrido completo del curso en bloques didacticos claros."
    )
    intro.paragraph_format.space_after = Pt(8)
    for section in SECTIONS:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(section["title"] + ": ").bold = True
        p.add_run(section["summary"])
    doc.add_heading("Videos Complementarios", level=1)
    vp = doc.add_paragraph(
        "Estos enlaces complementan las practicas y explicaciones del manual."
    )
    vp.paragraph_format.space_after = Pt(8)
    for links in VIDEO_LINKS.values():
        for label, url in links:
            p = doc.add_paragraph(style="List Bullet")
            add_hyperlink(p, label, url)
    doc.add_page_break()


def add_video_links_for_page(doc: Document, page_num: int) -> None:
    links = VIDEO_LINKS.get(page_num)
    if not links:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Recurso en video: ")
    r.bold = True
    r.font.color.rgb = RGBColor(85, 55, 31)
    for idx, (label, url) in enumerate(links):
        if idx:
            p.add_run(" | ")
        add_hyperlink(p, label, url)


def build_docx() -> Path:
    doc = Document()
    set_base_styles(doc)
    setup_section(doc.sections[0])
    add_cover(doc)
    add_module_map(doc)

    for sec_idx, section_data in enumerate(SECTIONS, 1):
        if not section_data["pages"]:
            continue
        doc.add_heading(section_data["title"], level=1)
        summary = doc.add_paragraph(section_data["summary"])
        summary.paragraph_format.space_after = Pt(10)
        sr = summary.runs[0]
        sr.italic = True
        sr.font.color.rgb = RGBColor(90, 90, 90)

        for page_num in section_data["pages"]:
            page_file = EXTRACT_DIR / f"page_{page_num:02d}.txt"
            raw = page_file.read_text(encoding="utf-8") if page_file.exists() else ""
            cleaned = clean_text(raw)
            page_images = iter_page_images(page_num)
            if not cleaned and not page_images:
                continue

            doc.add_heading(page_heading(page_num, cleaned or f"Pagina {page_num}"), level=2)
            if cleaned:
                for chunk in paragraph_chunks(cleaned):
                    chunk = re.sub(r"(?i)clic aqu[ií]", "Ver video complementario", chunk)
                    p = doc.add_paragraph(chunk)
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.line_spacing = 1.1

            add_video_links_for_page(doc, page_num)

            for idx, img_path in enumerate(page_images, 1):
                add_image_block(doc, img_path)

        if sec_idx != len(SECTIONS):
            doc.add_page_break()

    doc.save(FINAL_RAW_DOCX)
    write_prompt_catalog()
    return FINAL_RAW_DOCX


def write_prompt_catalog() -> None:
    lines = [
        "# Catalogo de prompts para regeneracion visual",
        "",
        "Estos prompts se prepararon para una ronda posterior de generacion de imagenes propias. En esta version del DOCX se aplico mejora tecnica local para mantener fidelidad visual.",
        "",
    ]
    for image_name, prompt in IMAGE_PROMPTS.items():
        lines.extend([f"## {image_name}", "", prompt, ""])
    PROMPTS_MD.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    path = build_docx()
    print(path)
